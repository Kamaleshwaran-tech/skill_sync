import io
import logging
from datetime import datetime, timedelta, timezone
from pathlib import Path
import uuid
import zipfile

import httpx
import pymupdf
import pytest
from docx import Document
from app.config.settings import get_settings
from app.database.session import SessionLocal
from app.models.models import Resume, ResumeAnalysis, ResumeMatchRun, Skill, UserSkill
from app.services.adzuna import AdzunaProvider, AdzunaError
from app.services.resume_parser import ResumeParser
from app.services.evidence import skill_evidence, experience_years
from app.services.job_matching import normalize_job, rank_jobs, score_job

API = "/api/v1"


def account(client):
    credentials = {
        "email": f"core-{uuid.uuid4().hex}@example.com",
        "password": "CoreTest123!",
    }
    user = client.post(
        API + "/auth/register", json={**credentials, "full_name": "Evidence Candidate"}
    )
    assert user.status_code == 200, user.text
    tokens = client.post(API + "/auth/login", json=credentials).json()
    return {"Authorization": "Bearer " + tokens["access_token"]}, user.json()


def document_bytes(text, table=False):
    document = Document()
    if table:
        document.add_table(rows=1, cols=1).cell(0, 0).text = text
    else:
        for line in text.splitlines():
            document.add_paragraph(line)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def upload(
    client,
    headers,
    text="Alex Sample\nSkills: Python, SQL\nSummary: I build tested web applications with readable code and documentation.",
):
    response = client.post(
        API + "/resumes/upload",
        headers=headers,
        files={
            "file": (
                "resume.docx",
                document_bytes(text),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 200, response.text
    assert "storage_path" not in response.json()
    resume_id = response.json()["id"]
    parsed = client.post(API + f"/resumes/{resume_id}/analyze", headers=headers)
    assert parsed.status_code == 200, parsed.text
    return resume_id


def raw_job(
    id="python-role",
    title="Python Developer",
    description="Required: Python, SQL. Preferred: Docker. Build database applications.",
):
    return {
        "id": id,
        "title": title,
        "company": {"display_name": "Synthetic Test Company"},
        "location": {"display_name": "Chennai"},
        "description": description,
        "created": datetime.now(timezone.utc).isoformat(),
        "redirect_url": "https://www.adzuna.in/jobs/land/ad/" + id,
        "salary_min": 500000,
        "salary_max": 900000,
        "salary_is_predicted": "1",
    }


@pytest.fixture
def adzuna(monkeypatch):
    """Mock only HTTP transport: real pagination, normalization, extraction, scoring and database writes execute."""
    settings = get_settings()
    monkeypatch.setattr(settings, "adzuna_app_id", "test-app-id")
    monkeypatch.setattr(settings, "adzuna_api_key", "test-key-must-never-leak")
    monkeypatch.setattr(settings, "job_max_retries", 0)
    requests = []
    pages = {1: {"results": [raw_job()], "count": 1}}
    real_client = httpx.Client

    def handle(request):
        requests.append(request)
        page = int(request.url.path.rsplit("/", 1)[1])
        payload = pages.get(page, {"results": [], "count": 0})
        if isinstance(payload, Exception):
            raise payload
        if isinstance(payload, int):
            return httpx.Response(payload, json={"error": "synthetic error"})
        return httpx.Response(200, json=payload)

    monkeypatch.setattr(
        "app.services.adzuna.httpx.Client",
        lambda **kwargs: real_client(transport=httpx.MockTransport(handle), **kwargs),
    )
    return pages, requests


def test_parser_never_invents_missing_location_education_or_experience():
    profile = ResumeParser().extract_structured(
        "Alex Sample\nSkills: Python and SQL\nI develop reliable applications with tests and readable documentation."
    )["profile"]
    assert profile["personal_info"]["location"] is None
    assert profile["education_text"] is None and profile["experience_years"] is None
    assert set(s["name"] for s in profile["technical_skills"]) == {"Python", "SQL"}
    assert "skills" in profile["sections"]
    assert "Perambalur" not in str(profile) and "B.E" not in str(profile)


def test_docx_tables_and_pdf_text_are_extracted(tmp_path):
    text = "Alex Sample\nSkills: Python, SQL\nEducation: Actual College\nDeveloping reliable systems with evidence and documentation."
    path = tmp_path / "table.docx"
    path.write_bytes(document_bytes(text, table=True))
    parsed = ResumeParser().extract(str(path))["profile"]
    assert any(s["name"] == "Python" for s in parsed["technical_skills"])
    path = tmp_path / "resume.pdf"
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_text((72, 72), text)
        pdf.save(path)
    assert (
        "Actual College"
        in ResumeParser().extract(str(path))["profile"]["education_text"]
    )


@pytest.mark.parametrize(
    "text,expected,absent",
    [
        (
            "JavaScript with Django and PostgreSQL",
            {"JavaScript", "Django", "PostgreSQL"},
            {"Java", "Go", "C"},
        ),
        ("Skills: C++ and C#", {"C++", "C#"}, {"C"}),
        ("We go further with Django applications", {"Django"}, {"Go"}),
    ],
)
def test_skill_boundaries(text, expected, absent):
    skills = set(skill_evidence(text))
    assert expected <= skills
    assert not skills & absent


def test_overlapping_experience_ranges_are_not_double_counted():
    years, _, method = experience_years(
        "Example", "Jan 2020 - Jan 2022\nJan 2021 - Jan 2023"
    )
    assert years == 3 and method == "estimated_from_month_ranges"
    assert experience_years("3 years of experience in Python")[0] == 3
    assert experience_years("A project completed for a client")[0] is None


def test_requirement_types_and_no_unsupported_score_bonus():
    job = normalize_job(
        raw_job(description="Required: Python, SQL. Preferred: Docker."), "in"
    )
    assert set(job["requirements"]["required"]) == {"Python", "SQL"}
    assert set(job["requirements"]["preferred"]) == {"Docker"}
    result = score_job(
        {"raw_text": "", "technical_skills": [], "experience_years": None}, job
    )
    assert result["match_score"] == 0
    assert "experience" not in {row["signal"] for row in result["score_breakdown"]}
    assert "education" not in str(result["score_breakdown"])


def test_ranking_uses_real_evidence_and_is_deterministic():
    profile = ResumeParser().extract_structured(
        "Alex Sample\nPython SQL developer with 3 years of experience building tested database applications."
    )["profile"]
    jobs = [
        raw_job(
            "wrong",
            "Java Developer",
            "Required: Java, Spring Boot. 5 years of experience.",
        ),
        raw_job(),
    ]
    ranked, skipped = rank_jobs(profile, jobs, "in", 30)
    assert ranked[0]["id"] == "python-role" and skipped == 0
    assert ranked == rank_jobs(profile, jobs, "in", 30)[0]
    assert 0 <= ranked[0]["match_score"] <= 100
    assert ranked[0]["comparisons"][0]["job_evidence"]
    assert (
        ranked[0]["salary_currency"] == "INR"
        and ranked[0]["salary_is_predicted"] is True
    )


def test_unsafe_url_and_html_removed_and_missing_metadata_preserved():
    raw = raw_job(description="<p>Required: Python.</p><script>alert(1)</script>")
    raw["redirect_url"] = "javascript:alert(1)"
    raw.pop("location")
    raw.pop("salary_min")
    raw.pop("salary_max")
    job = normalize_job(raw, "in")
    assert (
        job["application_url"] is None
        and "<script>" not in job["description"]
        and "alert" not in job["description"]
    )
    assert (
        job["location"] is None
        and job["salary_min"] is None
        and job["contract_time"] is None
    )


def test_missing_configuration_is_error_not_empty_matches(client):
    headers, _ = account(client)
    resume_id = upload(client, headers)
    status = client.get(API + "/jobs/status", headers=headers)
    assert status.json()["configured"] is False
    result = client.post(
        API + "/jobs/search", headers=headers, json={"resume_id": resume_id}
    )
    assert (
        result.status_code == 503
        and result.json()["detail"]["code"] == "provider_not_configured"
    )
    with SessionLocal() as db:
        assert db.query(ResumeMatchRun).count() == 0


def test_complete_pipeline_is_bound_to_selected_resume_not_account_skills(
    client, adzuna
):
    headers, user = account(client)
    first = upload(client, headers)
    second = upload(
        client,
        headers,
        "Alex Sample\nSkills: Java and Spring Boot\nBuilding reliable enterprise software applications with tests and documentation.",
    )
    # A manual account-level skill must not contaminate the selected resume.
    with SessionLocal() as db:
        skill = Skill(name="Docker", category="technical")
        db.add(skill)
        db.flush()
        db.add(UserSkill(user_id=user["id"], skill_id=skill.id, source="profile"))
        db.commit()
    first_result = client.post(
        API + "/jobs/search",
        headers=headers,
        json={
            "resume_id": first,
            "query": "Python developer",
            "location": "Chennai",
            "country": "in",
            "limit": 50,
        },
    )
    assert first_result.status_code == 200, first_result.text
    first_result = first_result.json()
    assert (
        first_result["resume_id"] == first
        and "Docker" in first_result["jobs"][0]["missing_skills"]
    )
    assert "Python" in first_result["jobs"][0]["matched_skills"]
    second_result = client.post(
        API + "/jobs/search",
        headers=headers,
        json={"resume_id": second, "query": "Python developer"},
    ).json()
    assert "Python" not in second_result["jobs"][0]["matched_skills"]
    assert (
        first_result["jobs"][0]["match_score"] > second_result["jobs"][0]["match_score"]
    )
    saved = client.get(API + f"/jobs/matches/{first}", headers=headers).json()["result"]
    assert (
        saved["from_saved_search"] is True
        and saved["analysis_id"] == first_result["analysis_id"]
    )
    assert saved["jobs"] == first_result["jobs"]
    request = adzuna[1][0]
    assert request.url.host == "api.adzuna.com"
    assert (
        request.url.params["what"] == "Python developer"
        and request.url.params["where"] == "Chennai"
    )
    assert "Alex" not in str(request.url) and "@" not in str(request.url)


def test_pagination_ranks_entire_fetched_set_and_filters_duplicates_expired(
    client, adzuna
):
    headers, _ = account(client)
    resume_id = upload(client, headers)
    pages, requests = adzuna
    pages[1] = {
        "results": [
            raw_job(str(i), "Java Developer", "Required: Java and Spring Boot.")
            for i in range(50)
        ],
        "count": 110,
    }
    expired = raw_job("expired")
    expired["created"] = (datetime.now(timezone.utc) - timedelta(days=60)).isoformat()
    pages[2] = {"results": [raw_job("best"), raw_job("0"), expired], "count": 110}
    result = client.post(
        API + "/jobs/search",
        headers=headers,
        json={"resume_id": resume_id, "limit": 100},
    ).json()
    assert result["jobs"][0]["id"] == "best"
    assert (
        result["scored_count"] == 51
        and result["fetched_count"] == 53
        and result["skipped_count"] == 2
    )
    assert result["provider_total"] == 110 and len(requests) == 2
    assert result["requested_limit"] == 100


@pytest.mark.parametrize(
    "status,code",
    [
        (401, "provider_authentication"),
        (403, "provider_authentication"),
        (429, "provider_rate_limited"),
        (500, "provider_error"),
    ],
)
def test_provider_http_errors_are_visible_and_keys_not_logged(
    client, adzuna, caplog, status, code
):
    headers, _ = account(client)
    resume_id = upload(client, headers)
    adzuna[0][1] = status
    with caplog.at_level(logging.INFO):
        response = client.post(
            API + "/jobs/search", headers=headers, json={"resume_id": resume_id}
        )
    assert response.status_code >= 500
    assert response.json()["detail"]["code"] == code
    assert (
        "test-key-must-never-leak" not in caplog.text and "app_key=" not in caplog.text
    )
    assert "test-key-must-never-leak" not in response.text
    with SessionLocal() as db:
        assert db.query(ResumeMatchRun).count() == 0


def test_partial_provider_failure_does_not_save_incomplete_success(client, adzuna):
    headers, _ = account(client)
    resume_id = upload(client, headers)
    adzuna[0][1] = {"results": [raw_job(str(i)) for i in range(50)], "count": 100}
    adzuna[0][2] = 503
    response = client.post(
        API + "/jobs/search",
        headers=headers,
        json={"resume_id": resume_id, "limit": 100},
    )
    assert response.status_code == 502
    with SessionLocal() as db:
        assert db.query(ResumeMatchRun).count() == 0


@pytest.mark.parametrize(
    "response,code",
    [
        ({"results": "bad"}, "provider_invalid_response"),
        (httpx.ReadTimeout("secret request URL"), "provider_timeout"),
    ],
)
def test_invalid_and_timeout_provider_responses(client, adzuna, response, code):
    headers, _ = account(client)
    resume_id = upload(client, headers)
    adzuna[0][1] = response
    result = client.post(
        API + "/jobs/search", headers=headers, json={"resume_id": resume_id}
    )
    assert result.status_code >= 500 and result.json()["detail"]["code"] == code


def test_valid_empty_search_is_not_an_error(client, adzuna):
    headers, _ = account(client)
    resume_id = upload(client, headers)
    adzuna[0][1] = {"results": [], "count": 0}
    result = client.post(
        API + "/jobs/search", headers=headers, json={"resume_id": resume_id}
    )
    assert (
        result.status_code == 200
        and result.json()["jobs"] == []
        and result.json()["scored_count"] == 0
    )


def test_cross_account_ownership_and_protected_routes(client, adzuna):
    headers, _ = account(client)
    other, _ = account(client)
    resume_id = upload(client, headers)
    assert (
        client.post(
            API + "/jobs/search", headers=headers, json={"resume_id": resume_id}
        ).status_code
        == 200
    )
    for path in [
        f"/resumes/{resume_id}/file",
        f"/resumes/{resume_id}/analysis",
        f"/jobs/matches/{resume_id}",
    ]:
        assert client.get(API + path, headers=other).status_code == 404
        assert client.get(API + path).status_code == 401
    assert (
        client.post(
            API + "/jobs/search", headers=other, json={"resume_id": resume_id}
        ).status_code
        == 404
    )
    assert (
        client.delete(API + f"/resumes/{resume_id}", headers=other).status_code == 404
    )
    assert client.get(API + "/jobs/status").status_code == 401


def test_idempotent_analysis_and_delete_remove_file_and_matches(client, adzuna):
    headers, _ = account(client)
    resume_id = upload(client, headers)
    assert (
        client.post(API + f"/resumes/{resume_id}/analyze", headers=headers).status_code
        == 200
    )
    assert (
        client.post(
            API + "/jobs/search", headers=headers, json={"resume_id": resume_id}
        ).status_code
        == 200
    )
    with SessionLocal() as db:
        assert db.query(ResumeAnalysis).filter_by(resume_id=resume_id).count() == 1
        path = Path(db.get(Resume, resume_id).storage_path)
        assert path.exists()
    assert (
        client.delete(API + f"/resumes/{resume_id}", headers=headers).status_code == 200
    )
    assert not path.exists()
    with SessionLocal() as db:
        assert (
            db.query(ResumeMatchRun).count() == 0
            and db.query(ResumeAnalysis).count() == 0
        )


def test_old_analysis_is_reparsed_without_invented_data(client, adzuna):
    headers, _ = account(client)
    resume_id = upload(client, headers)
    with SessionLocal() as db:
        old = db.query(ResumeAnalysis).filter_by(resume_id=resume_id).first()
        old.profile = {"name": "Invented", "technical_skills": [{"name": "Docker"}]}
        db.commit()
    assert (
        client.post(
            API + "/jobs/search", headers=headers, json={"resume_id": resume_id}
        ).status_code
        == 409
    )
    assert (
        client.post(API + f"/resumes/{resume_id}/analyze", headers=headers).status_code
        == 200
    )
    assert (
        client.post(
            API + "/jobs/search", headers=headers, json={"resume_id": resume_id}
        ).status_code
        == 200
    )


@pytest.mark.parametrize(
    "payload",
    [
        {"country": "xx"},
        {"limit": 101},
        {"limit": 75},
        {"max_days": 0},
        {"query": "x" * 121},
        {"resume_text": "private"},
    ],
)
def test_search_filters_are_bounded(client, payload):
    headers, _ = account(client)
    assert (
        client.post(
            API + "/jobs/search", headers=headers, json={"resume_id": 1, **payload}
        ).status_code
        == 422
    )


@pytest.mark.parametrize(
    "filename,data",
    [
        ("../private.docx", b"123"),
        ("bad.pdf", b"not PDF"),
        ("bad.docx", b"not zip"),
        ("test.exe", b"MZ"),
        ("empty.pdf", b""),
    ],
)
def test_upload_rejects_bad_files(client, filename, data):
    headers, _ = account(client)
    assert (
        client.post(
            API + "/resumes/upload", headers=headers, files={"file": (filename, data)}
        ).status_code
        == 400
    )


def test_expanded_docx_limit(client):
    headers, _ = account(client)
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "x")
        archive.writestr("word/document.xml", "x" * (26 * 1024 * 1024))
    assert (
        client.post(
            API + "/resumes/upload",
            headers=headers,
            files={"file": ("bomb.docx", stream.getvalue())},
        ).status_code
        == 400
    )


def test_out_of_scope_routes_removed(client):
    paths = client.get(API + "/openapi.json").json()["paths"]
    for term in [
        "roadmap",
        "career-readiness",
        "courses",
        "candidates",
        "admin",
        "applications",
        "recommendations",
        "profile",
    ]:
        assert not any(term in path for path in paths)
    assert (
        client.post(
            API + "/auth/register",
            json={
                "email": "employer@example.com",
                "password": "CoreTest123!",
                "role": "EMPLOYER",
            },
        ).status_code
        == 422
    )


@pytest.mark.parametrize("lease_age", [None, 6])
def test_interrupted_parsing_can_be_retried(client, lease_age):
    headers, _ = account(client)
    resume_id = upload(client, headers)
    with SessionLocal() as db:
        resume = db.get(Resume, resume_id)
        resume.status = "PROCESSING"
        resume.processing_started_at = (
            None
            if lease_age is None
            else datetime.now(timezone.utc).replace(tzinfo=None)
            - timedelta(minutes=lease_age)
        )
        db.commit()
    response = client.post(API + f"/resumes/{resume_id}/analyze", headers=headers)
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "COMPLETED"


def test_active_parse_is_locked_but_stale_parse_can_be_deleted(client):
    headers, _ = account(client)
    resume_id = upload(client, headers)
    with SessionLocal() as db:
        resume = db.get(Resume, resume_id)
        resume.status = "PROCESSING"
        resume.processing_started_at = datetime.now(timezone.utc).replace(tzinfo=None)
        db.commit()
    assert (
        client.post(API + f"/resumes/{resume_id}/analyze", headers=headers).status_code
        == 409
    )
    assert (
        client.delete(API + f"/resumes/{resume_id}", headers=headers).status_code == 409
    )
    with SessionLocal() as db:
        db.get(Resume, resume_id).processing_started_at = datetime.now(
            timezone.utc
        ).replace(tzinfo=None) - timedelta(minutes=6)
        db.commit()
    assert (
        client.delete(API + f"/resumes/{resume_id}", headers=headers).status_code == 200
    )


@pytest.mark.parametrize(
    "text",
    [
        "No experience with Python. Skills: Java.",
        "I have not used Python. Skills: Java.",
        "Python experience is not required. Java required.",
    ],
)
def test_explicit_negative_skill_mentions_are_not_positive_evidence(text):
    assert "Python" not in skill_evidence(text)
    assert "Java" in skill_evidence(text)


def test_mixed_required_preferred_clauses_and_experience_ranges():
    from app.services.job_matching import extract_requirements
    from app.services.evidence import job_experience

    groups = extract_requirements(
        "Developer", "Python and SQL required, Docker preferred. Java is not required."
    )
    assert set(groups["required"]) == {"Python", "SQL"}
    assert set(groups["preferred"]) == {"Docker"}
    assert "Java" not in groups["mentioned"]
    assert job_experience("3 to 5 years of experience") == 3
    assert experience_years("3–5 years of experience")[0] == 3


def test_superseded_parser_cannot_overwrite_newer_lease(client, monkeypatch):
    headers, _ = account(client)
    resume_id = upload(client, headers)
    with SessionLocal() as db:
        db.get(Resume, resume_id).status = "FAILED"
        db.commit()
        count = db.query(ResumeAnalysis).filter_by(resume_id=resume_id).count()
    original = ResumeParser.extract

    def replace_lease(self, path):
        result = original(self, path)
        with SessionLocal() as db:
            db.get(Resume, resume_id).processing_started_at = datetime.now(
                timezone.utc
            ).replace(tzinfo=None) + timedelta(seconds=1)
            db.commit()
        return result

    monkeypatch.setattr(ResumeParser, "extract", replace_lease)
    assert (
        client.post(API + f"/resumes/{resume_id}/analyze", headers=headers).status_code
        == 409
    )
    with SessionLocal() as db:
        assert db.get(Resume, resume_id).status == "PROCESSING"
        assert db.query(ResumeAnalysis).filter_by(resume_id=resume_id).count() == count


def test_phone_search_skips_invalid_short_candidate():
    profile = ResumeParser().extract_structured(
        "Alex Sample\nReference 123456789\nPhone: +1 202 555 0123\nSkills: Python and SQL development with tests."
    )["profile"]
    assert profile["personal_info"]["phone"] == "+1 202 555 0123"


def test_analysis_changed_while_provider_fetching_is_not_saved(client, monkeypatch):
    headers, _ = account(client)
    resume_id = upload(client, headers)

    def fetch(*args):
        with SessionLocal() as db:
            old = db.query(ResumeAnalysis).filter_by(resume_id=resume_id).first()
            db.add(ResumeAnalysis(resume_id=resume_id, profile=old.profile))
            db.commit()
        return {"results": [raw_job()], "source": "test_fixture", "provider_count": 1}

    monkeypatch.setattr(AdzunaProvider, "fetch_jobs", fetch)
    response = client.post(
        API + "/jobs/search", headers=headers, json={"resume_id": resume_id}
    )
    assert response.status_code == 409, response.text
    with SessionLocal() as db:
        assert db.query(ResumeMatchRun).filter_by(resume_id=resume_id).count() == 0


def test_configuration_validation_does_not_echo_provider_secrets():
    from app.config.settings import Settings
    from pydantic import ValidationError

    with pytest.raises(ValidationError) as failure:
        Settings(
            jwt_secret_key="short",
            adzuna_api_key="secret-redaction-check",
            _env_file=None,
        )
    assert "secret-redaction-check" not in str(failure.value)
    assert "input_value" not in str(failure.value)


@pytest.mark.parametrize(
    "secret", ["skillsync-super-secure-jwt-secret-key-32-chars-long", "a" * 64]
)
def test_public_example_and_repeated_jwt_secrets_are_rejected(secret):
    from app.config.settings import Settings
    from pydantic import ValidationError

    with pytest.raises(ValidationError):
        Settings(jwt_secret_key=secret, _env_file=None)
