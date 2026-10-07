import os
import io
import tempfile
import pytest

from app.services.resume_parser import ResumeParser


def test_extract_structured_from_text():
    parser = ResumeParser()
    sample = """
    John Doe
    Email: john.doe@example.com
    Phone: +1 555-123-4567

    Summary
    Experienced Full Stack Developer with React, Node.js, and PostgreSQL experience.

    Education
    B.S. Computer Science, University of Example, 2016 - 2020

    Experience
    Software Engineer at Acme Corp (2020 - Present)
    Worked with React, TypeScript, Docker, AWS and Kubernetes.

    Skills
    React, Node.js, TypeScript, Docker, AWS, PostgreSQL, Python
    """

    res = parser.extract_structured(sample)
    assert res and 'profile' in res
    profile = res['profile']
    assert profile['name'] is not None
    assert any('john.doe@example.com' in e for e in profile['emails'])
    tech = profile['technical_skills']
    # expect React and Node.js normalized
    names = [s['name'] for s in tech]
    assert 'React' in names or 'Node.js' in names


@pytest.mark.skipif(not os.environ.get('RUN_DOCX_TESTS'), reason="DOCX tests disabled")
def test_extract_from_docx_file(tmp_path):
    # create a real docx if python-docx present
    try:
        import docx
    except Exception:
        pytest.skip('python-docx not installed')

    doc = docx.Document()
    doc.add_paragraph('Jane Smith')
    doc.add_paragraph('Email: jane.smith@example.com')
    doc.add_paragraph('Skills')
    doc.add_paragraph('ReactJS, Python, Docker, AWS')
    p = tmp_path / "resume.docx"
    doc.save(str(p))

    parser = ResumeParser()
    text = parser.extract_text_from_docx(str(p))
    res = parser.extract_structured(text)
    assert res and res['profile']['name'] is not None
    assert any('jane.smith@example.com' in e for e in res['profile']['emails'])
    names = [s['name'] for s in res['profile']['technical_skills']]
    assert 'React' in names or 'Python' in names
