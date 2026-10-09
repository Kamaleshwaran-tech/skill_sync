"""Browser performance assertions against tests.serve_browser_fixture (never live jobs).
Start fixture API 8001 and frontend proxy 5174. Requires playwright + python-docx.
Uses synthetic accounts/resumes; prints timings, never credentials or tokens.
"""

import asyncio
import io
import json
import os
import statistics
import time
import uuid
from docx import Document
from playwright.async_api import async_playwright, expect

BASE = os.environ.get("BASE_URL", "http://127.0.0.1:5174")


async def main():
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True, args=["--no-sandbox"])
        context = await browser.new_context(viewport={"width": 1440, "height": 1000})
        api = context.request

        async def account():
            credentials = {
                "email": f"perf-{uuid.uuid4().hex}@example.com",
                "password": "PerformanceOnly123!",
            }
            response = await api.post(
                BASE + "/api/v1/auth/register",
                data={**credentials, "full_name": "Synthetic Performance Test"},
            )
            assert response.ok, await response.text()
            response = await api.post(BASE + "/api/v1/auth/login", data=credentials)
            return credentials, await response.json()

        _, tokens = await account()
        headers = {"Authorization": "Bearer " + tokens["access_token"]}

        async def resume(filename, name, skills, query, limit):
            doc = Document()
            doc.add_paragraph(name)
            doc.add_paragraph("Skills: " + skills)
            doc.add_paragraph(
                "I build tested software applications with readable documentation and reliable releases."
            )
            buf = io.BytesIO()
            doc.save(buf)
            r = await api.post(
                BASE + "/api/v1/resumes/upload",
                headers=headers,
                multipart={
                    "file": {
                        "name": filename,
                        "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        "buffer": buf.getvalue(),
                    }
                },
            )
            assert r.ok, await r.text()
            rid = (await r.json())["id"]
            r = await api.post(BASE + f"/api/v1/resumes/{rid}/analyze", headers=headers)
            assert r.ok
            r = await api.post(
                BASE + "/api/v1/jobs/search",
                headers=headers,
                data={"resume_id": rid, "query": query, "limit": limit},
            )
            assert r.ok
            assert (await r.json())["provider"] == "test_fixture"
            return str(rid)

        first = await resume(
            "perf-python.docx", "Performance Python", "Python, SQL", "many-jobs", 100
        )
        second = await resume(
            "perf-java.docx",
            "Performance Java",
            "Java, Spring Boot",
            "Java developer",
            20,
        )
        page = await context.new_page()
        await page.goto(BASE)
        await page.evaluate(
            '(tokens)=>{localStorage.setItem("skillsync.access_token",tokens.access_token);localStorage.setItem("skillsync.refresh_token",tokens.refresh_token)}',
            tokens,
        )
        requests = []
        errors = []
        saved_done = asyncio.Event()
        status_done = asyncio.Event()
        delayed_saved_finished = asyncio.Event()
        page.on(
            "request",
            lambda r: (
                requests.append((r.method, r.url)) if "/api/v1/" in r.url else None
            ),
        )
        page.on("pageerror", lambda e: errors.append(str(e)))
        page.on(
            "response",
            lambda r: saved_done.set() if f"/jobs/matches/{second}" in r.url else None,
        )

        page.on(
            "response",
            lambda r: status_done.set() if r.url.endswith("/jobs/status") else None,
        )

        async def slow_status(route):
            await asyncio.sleep(1.2)
            await route.continue_()

        await page.route("**/jobs/status", slow_status)

        async def slow_saved(route):
            # Capture the old database snapshot first, then delay its delivery.
            response = await route.fetch()
            payload = await response.json()
            await asyncio.sleep(0.9)
            try:
                await route.fulfill(
                    status=response.status,
                    content_type="application/json",
                    body=json.dumps(payload),
                )
            finally:
                delayed_saved_finished.set()

        await page.route("**/jobs/matches/*", slow_saved)
        await page.reload()
        await expect(
            page.get_by_role("heading", name="Performance Java", exact=True)
        ).to_be_visible()
        assert (
            not saved_done.is_set()
        ), "Analysis unnecessarily waited for the delayed saved-jobs request"
        assert not status_done.is_set(), "Analysis waited for provider status"
        await asyncio.wait_for(saved_done.wait(), 10)
        await expect(page.locator("article")).to_have_count(2)
        await page.get_by_label("Selected resume", exact=True).select_option(first)
        await expect(page.locator("article")).to_have_count(20)
        assert (
            await page.locator(".comparison-list").count() == 0
        ), "Closed evidence rendered eagerly"
        await page.locator("article").first.locator("summary").click()
        await expect(page.locator(".comparison-list")).to_have_count(1)
        await page.locator("article").first.locator("summary").click()
        await expect(page.locator(".comparison-list")).to_have_count(0)
        before = len(requests)
        await page.get_by_role("button", name="Show next 20 jobs", exact=True).click()
        await expect(page.locator("article")).to_have_count(40)
        assert len(requests) == before, "Show more unexpectedly fetched jobs again"
        await page.locator("article").first.evaluate(
            'node=>node.dataset.retained="yes"'
        )
        for _ in range(3):
            await page.get_by_role("button", name="Resume analysis", exact=True).click()
            await expect(
                page.get_by_role("heading", name="Extracted personal information")
            ).to_be_visible()
            await page.get_by_role("button", name="Job matching", exact=True).click()
        assert len(requests) == before, "Page navigation refetched data"
        assert (
            await page.locator("article").first.get_attribute("data-retained") == "yes"
        ), "Cards remounted on tab navigation"
        await page.get_by_role("link", name="SkillSync home", exact=True).click()
        assert len(requests) == before, "Home navigation caused a document reload"
        warm = []
        for _ in range(3):
            for rid, name, count in [
                (second, "Performance Java", 2),
                (first, "Performance Python", 20),
            ]:
                start = time.perf_counter()
                await page.get_by_label("Selected resume", exact=True).select_option(
                    rid
                )
                await expect(
                    page.get_by_role("heading", name=name, exact=True)
                ).to_be_visible()
                await expect(page.locator("article")).to_have_count(count)
                warm.append((time.perf_counter() - start) * 1000)
        assert len(requests) == before, "Warm resume switching refetched cached GETs"
        await expect(
            page.get_by_text(
                "These are saved results, not a live availability check.", exact=False
            )
        ).to_be_visible()
        # Race: a successful fresh POST must beat an older saved GET still in flight.
        await page.unroute("**/jobs/status", slow_status)
        delayed_saved_finished.clear()
        await page.reload()
        await expect(
            page.get_by_role("heading", name="Performance Java", exact=True)
        ).to_be_visible()
        assert not delayed_saved_finished.is_set()
        await page.get_by_label("Job title or keywords", exact=True).fill(
            "new-performance-query"
        )
        async with page.expect_response(lambda r: r.url.endswith("/jobs/search")):
            await page.get_by_role("button", name="Find matches", exact=False).click()
        await asyncio.wait_for(delayed_saved_finished.wait(), 10)
        await expect(
            page.get_by_text("Results for “new-performance-query”", exact=False)
        ).to_be_visible()
        await page.get_by_label("Selected resume", exact=True).select_option(first)
        await expect(page.locator("article")).to_have_count(20)
        # A new live search must still perform POST, retaining the last snapshot while waiting.
        started = asyncio.Event()
        release = asyncio.Event()

        async def hold_search(route):
            started.set()
            await release.wait()
            await route.continue_()

        await page.route("**/jobs/search", hold_search)
        async with page.expect_response(
            lambda r: r.url.endswith("/jobs/search")
        ) as fresh:
            await page.get_by_role("button", name="Find matches", exact=False).click()
            await asyncio.wait_for(started.wait(), 10)
            await expect(page.locator("article")).to_have_count(20)
            await page.get_by_role("button", name="Resume analysis", exact=True).click()
            await expect(
                page.get_by_role("heading", name="Extracted personal information")
            ).to_be_visible()
            release.set()
        assert (await fresh.value).status == 200
        await page.unroute("**/jobs/search", hold_search)
        await page.get_by_role("button", name="Job matching", exact=True).click()
        async with page.expect_response(
            lambda r: r.url.endswith("/jobs/search")
        ) as fresh_again:
            await page.get_by_role("button", name="Find matches", exact=False).click()
        assert (await fresh_again.value).status == 200
        # Manual invalidation deliberately refetches the selected resume's data.
        matches_before = sum("/jobs/matches/" + first in url for _, url in requests)
        async with page.expect_response(lambda r: "/jobs/matches/" + first in r.url):
            await page.get_by_role(
                "button", name="Refresh saved data", exact=True
            ).click()
        assert (
            sum("/jobs/matches/" + first in url for _, url in requests)
            == matches_before + 1
        )
        await expect(
            page.get_by_role("button", name="Find matches", exact=False)
        ).to_be_enabled()
        new_credentials, _ = await account()
        await page.get_by_role("button", name="Sign out", exact=True).click()
        await expect(page.get_by_role("heading", name="Welcome back")).to_be_visible()
        await page.get_by_label("Email address", exact=True).fill(
            new_credentials["email"]
        )
        await page.get_by_label("Password", exact=True).fill(
            new_credentials["password"]
        )
        await page.get_by_role("button", name="Sign in", exact=False).first.click()
        await expect(
            page.get_by_role("button", name="Sign out", exact=True)
        ).to_be_visible()
        await expect(page.locator("article")).to_have_count(0)
        await expect(page.get_by_label("Selected resume", exact=True)).to_have_count(0)
        assert not errors, errors
        print(
            json.dumps(
                {
                    "result": "passed",
                    "provider": "Synthetic fixtures only",
                    "saved_request_delay_ms": 900,
                    "provider_status_delay_ms": 1200,
                    "analysis_visible_before_provider_status": True,
                    "analysis_visible_before_saved_jobs": True,
                    "warm_navigation_extra_requests": 0,
                    "warm_resume_switch_median_ms": round(statistics.median(warm), 1),
                    "warm_resume_switch_max_ms": round(max(warm), 1),
                    "initial_cards_for_100_results": 20,
                    "closed_explanation_nodes": 0,
                    "retained_cards_during_refresh": True,
                    "fresh_search_always_uses_network": True,
                    "late_saved_response_does_not_replace_fresh_search": True,
                    "manual_refresh_invalidates": True,
                    "new_account_does_not_receive_old_snapshots": True,
                    "page_errors": errors,
                },
                indent=2,
            )
        )
        await browser.close()


if __name__ == "__main__":
    asyncio.run(main())
