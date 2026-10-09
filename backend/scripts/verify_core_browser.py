"""Chromium integration test for the labelled tests.serve_browser_fixture server.

Run fixture API on 8001; frontend proxy on 5174 with VITE_BACKEND_URL=http://127.0.0.1:8001.
BASE_URL=http://127.0.0.1:5174 python scripts/verify_core_browser.py
Requires optional playwright + python-docx. Never submits real job applications.
"""

import io
import json
import os
import uuid
from pathlib import Path
from docx import Document
from playwright.sync_api import sync_playwright, expect

base = os.environ.get("BASE_URL", "http://127.0.0.1:5174")
checks = []
with sync_playwright() as p:
    browser = p.chromium.launch(headless=True, args=["--no-sandbox"])
    context = browser.new_context(viewport={"width": 1440, "height": 1000})
    page = context.new_page()
    errors = []
    page.on("pageerror", lambda error: errors.append(str(error)))
    page.goto(base)
    page.get_by_role("button", name="Create account", exact=True).click()
    page.get_by_label("Full name", exact=True).fill("Browser Evidence")
    email = f"browser-{uuid.uuid4().hex}@example.com"
    page.get_by_label("Email address", exact=True).fill(email)
    page.get_by_label("Password", exact=True).fill("BrowserCore123!")
    page.get_by_role("button", name="Create account", exact=False).click()
    expect(page.get_by_text("Account created.", exact=False)).to_be_visible()
    page.get_by_label("Email address", exact=True).fill(email)
    page.get_by_label("Password", exact=True).fill("BrowserCore123!")
    page.get_by_role("button", name="Sign in", exact=False).first.click()
    expect(page.get_by_text("TEST FIXTURES —", exact=False)).to_be_visible()
    checks.append("Registration and login through real API")

    def upload(skills, name):
        document = Document()
        document.add_paragraph("Browser Evidence")
        document.add_paragraph("Skills: " + skills)
        document.add_paragraph(
            "I build tested software applications with readable documentation and reliable releases."
        )
        stream = io.BytesIO()
        document.save(stream)
        with page.expect_response(
            lambda r: r.url.endswith("/jobs/search") and r.request.method == "POST",
            timeout=30000,
        ) as matched:
            page.locator("input[type=file]").set_input_files(
                {
                    "name": name,
                    "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    "buffer": stream.getvalue(),
                }
            )
        assert matched.value.status == 200, matched.value.text()
        assert matched.value.json()["provider"] == "test_fixture"

    upload("Python, SQL", "python-resume.docx")
    expect(
        page.locator("article").first.get_by_role(
            "heading", name="TEST FIXTURE — Python Developer"
        )
    ).to_be_visible()
    page.locator("article").first.locator("summary").click()
    expect(
        page.locator("article")
        .first.get_by_text("Skills: Python, SQL", exact=False)
        .first
    ).to_be_visible()
    checks.extend(
        [
            "DOCX extraction",
            "Automatic Adzuna-shaped HTTP fixture fetch",
            "Evidence-based ranking and displayed explanation",
        ]
    )
    first_card = page.locator("article").first
    expect(
        first_card.get_by_role("region", name="Missing skills", exact=True)
    ).to_contain_text("Docker")
    details_button = first_card.get_by_role("button", name="Details", exact=True)
    details_button.click()
    dialog = page.get_by_role(
        "dialog", name="TEST FIXTURE — Python Developer", exact=True
    )
    expect(dialog).to_be_visible()
    expect(dialog.get_by_text("Salary", exact=True)).to_be_visible()
    expect(
        dialog.get_by_role("region", name="Missing skills", exact=True)
    ).to_contain_text("Docker")
    expect(dialog).to_contain_text("python-resume.docx")
    for _ in range(5):
        page.keyboard.press("Tab")
        assert dialog.evaluate("node => node.contains(document.activeElement)"), (
            "Focus escaped modal"
        )
    for _ in range(4):
        page.keyboard.press("Shift+Tab")
        assert dialog.evaluate("node => node.contains(document.activeElement)"), (
            "Backward focus escaped modal"
        )
    page.keyboard.press("Escape")
    expect(dialog).not_to_be_visible()
    expect(details_button).to_be_focused()
    details_button.press("Enter")
    expect(dialog).to_be_visible()
    dialog.get_by_role("button", name="Close job details", exact=True).click()
    expect(dialog).not_to_be_visible()
    page.locator("article").nth(1).get_by_role(
        "button", name="Details", exact=True
    ).click()
    second_dialog = page.get_by_role(
        "dialog", name="TEST FIXTURE — Java Developer", exact=True
    )
    expect(second_dialog).to_be_visible()
    expect(
        second_dialog.get_by_role("region", name="Missing skills", exact=True)
    ).to_contain_text("Java")
    second_dialog.get_by_role("button", name="Close", exact=True).click()
    page.set_viewport_size({"width": 390, "height": 844})
    details_button.click()
    expect(dialog).to_be_visible()
    assert dialog.evaluate("node => node.scrollWidth <= node.clientWidth + 1"), (
        "Dialog mobile overflow"
    )
    assert page.evaluate(
        "document.documentElement.scrollWidth <= window.innerWidth+2"
    ), "Mobile page overflow"
    screenshot = Path.home() / ".cache" / "skillsync-job-details-mobile.png"
    screenshot.parent.mkdir(parents=True, exist_ok=True)
    page.screenshot(path=str(screenshot))
    page.mouse.click(2, 2)
    expect(dialog).not_to_be_visible()
    assert page.evaluate("document.body.style.overflow !== 'hidden'")
    page.set_viewport_size({"width": 1440, "height": 1000})
    checks.extend(
        [
            "Details opens the correct job popup on every card",
            "Visible missing skills on cards and popup",
            "Modal keyboard focus, Escape, close button and backdrop dismissal",
            "Mobile popup fits without horizontal overflow",
        ]
    )
    page.reload()
    expect(
        page.get_by_text(
            "These are saved results, not a live availability check.", exact=False
        )
    ).to_be_visible()
    expect(
        page.locator("article").first.get_by_role(
            "heading", name="TEST FIXTURE — Python Developer"
        )
    ).to_be_visible()
    checks.append("Saved matching result after reload")
    page.evaluate("localStorage.setItem('skillsync.access_token','expired-test-token')")
    with page.expect_response(lambda r: r.url.endswith("/auth/refresh")) as rotated:
        page.reload()
    assert rotated.value.status == 200
    expect(page.get_by_role("button", name="Sign out", exact=True)).to_be_visible()
    checks.append("Access token refresh")
    upload("Java, Spring Boot", "java-resume.docx")
    expect(
        page.locator("article").first.get_by_role(
            "heading", name="TEST FIXTURE — Java Developer"
        )
    ).to_be_visible()
    expect(
        page.locator("article").first.get_by_role(
            "region", name="Missing skills", exact=True
        )
    ).to_contain_text("None among the skills identified")
    page.get_by_label("Selected resume", exact=True).select_option(
        label="python-resume.docx · completed"
    )
    expect(
        page.locator("article").first.get_by_role(
            "heading", name="TEST FIXTURE — Python Developer"
        )
    ).to_be_visible()
    checks.append("Two resumes remain isolated; switching restores the correct matches")
    expect(
        page.locator("article").first.get_by_role(
            "region", name="Missing skills", exact=True
        )
    ).to_contain_text("Docker")
    page.get_by_label("Job title or keywords", exact=True).fill("provider-fail")
    page.get_by_role("button", name="Find matches", exact=False).click()
    expect(page.get_by_role("alert")).to_contain_text("Adzuna rejected")
    expect(page.locator("article")).to_have_count(0)
    page.get_by_label("Job title or keywords", exact=True).fill("no-results")
    page.get_by_role("button", name="Find matches", exact=False).click()
    expect(
        page.get_by_role("heading", name="No current results for this search")
    ).to_be_visible()
    checks.append("Provider failure differs from a legitimate empty result")
    page.get_by_label("Job title or keywords", exact=True).fill("Python developer")
    page.get_by_role("button", name="Find matches", exact=False).click()
    expect(page.locator("article").first).to_be_visible()
    page.set_viewport_size({"width": 390, "height": 844})
    assert page.evaluate(
        "document.documentElement.scrollWidth <= window.innerWidth+2"
    ), "Mobile overflow"
    checks.append("Mobile layout at 390px")
    page.set_viewport_size({"width": 1440, "height": 1000})
    screenshot = Path(
        os.environ.get(
            "SCREENSHOT", str(Path.home() / ".cache" / "skillsync-core-workspace.png")
        )
    )
    screenshot.parent.mkdir(parents=True, exist_ok=True)
    page.screenshot(path=str(screenshot), full_page=True)
    page.on("dialog", lambda dialog: dialog.accept())
    page.get_by_role("button", name="Delete resume", exact=True).click()
    expect(
        page.get_by_label("Selected resume", exact=True).locator("option")
    ).to_have_count(1)
    expect(
        page.locator("article").first.get_by_role(
            "heading", name="TEST FIXTURE — Java Developer"
        )
    ).to_be_visible()
    checks.append("Resume deletion switches to remaining resume, not deleted data")
    page.get_by_role("button", name="Sign out", exact=True).click()
    expect(page.get_by_role("heading", name="Welcome back")).to_be_visible()
    assert page.evaluate("localStorage.getItem('skillsync.access_token')") is None
    assert not errors, errors
    checks.append("Logout clears the session")
    print(
        json.dumps(
            {
                "result": "passed",
                "provider": "TEST FIXTURES — not a live Adzuna verification",
                "checks": checks,
                "pageErrors": errors,
            },
            indent=2,
        )
    )
    browser.close()
